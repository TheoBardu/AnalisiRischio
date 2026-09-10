#! /usr/bin/env python3
"""
Lettura dei risultati della valutazione del rischio vibrazioni.

Fonte primaria: riepilogo_vibrazioni.json, che il runner scrive in log/
(accanto a output/, nel ramo Vibrazioni) subito dopo l'analisi a partire da analisi.riepilogo(). E' il modo piu' fedele,
perche' sono gli stessi numeri calcolati dal backend.

Ripiego: VR_VIB.xlsx, quando in cartella c'e' il risultato di un'esecuzione
precedente ma non il json. I fogli 'Scheda N' hanno una geometria variabile
(la tabella viene ridimensionata al numero di attrezzature), quindi invece di
contare le righe ci si ancora all'etichetta 'VALUTAZIONE SU BASE GIORNALIERA'.
"""

import json
import os

import openpyxl

NOME_RIEPILOGO_JSON = 'riepilogo_vibrazioni.json'
PREFISSO_SCHEDA = 'scheda'

ETICHETTA_VALUTAZIONE = 'VALUTAZIONE SU BASE GIORNALIERA'

# offset rispetto alla riga dell'etichetta di valutazione
OFFSET_A8 = 1
OFFSET_UA8 = 2
OFFSET_CLASSE = 3

RIGA_INIZIO_DATI = 9
COL_VALORE_INTESTAZIONE = 4
RIGHE_INTESTAZIONE = {2: 'ID_GrOm', 3: 'Descrizione_GrOm',
                      4: 'ID_reparto', 5: 'Descrizione_reparto'}

COL_HAV = {'ID': 2, 'Categoria': 3, 'Marca': 4, 'Modello': 5,
           'Aw': 6, 'U': 7, 'Te': 8}
COL_WBV = {'ID': 12, 'Categoria': 13, 'Marca': 14, 'Modello': 15,
           'Aw': 16, 'U': 17, 'Te': 18}

COL_CLASSE = {'HAV': 2, 'WBV': 12}
COL_A8 = {'HAV': 4, 'WBV': 14}
COL_UA8_AFFIANCO = {'HAV': 6, 'WBV': 16}
COL_A8MAX = {'HAV': 8, 'WBV': 18}
COL_DY_HAV = 8

COLORI_CLASSE = {'BASSA': '#32cd32', 'MEDIA': '#00bfff', 'ALTA': '#b22222'}


def _testo(valore):
    if valore is None:
        return ''
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def _numero(valore):
    if valore is None or valore == '':
        return None
    try:
        return float(str(valore).replace(',', '.'))
    except (TypeError, ValueError):
        return None


def _fmt(valore, decimali=2):
    numero = _numero(valore)
    return '—' if numero is None else f'{numero:.{decimali}f}'


def _classe_breve(testo):
    """'Classe di rischio alta' -> 'ALTA'."""
    basso = (testo or '').strip().lower()
    for chiave in ('bassa', 'media', 'alta'):
        if basso.endswith(chiave):
            return chiave.upper()
    return ''


def _classe_da_valore(a8max, tipo, parametri):
    """Classe calcolata dalle soglie, quando il foglio non la riporta."""
    valore = _numero(a8max)
    if valore is None:
        return ''
    media = parametri.get(f'SOGLIA_MEDIA_{tipo}', 2.5 if tipo == 'HAV' else 0.5)
    alta = parametri.get(f'SOGLIA_ALTA_{tipo}', 5.0 if tipo == 'HAV' else 1.0)
    if valore >= alta:
        return 'ALTA'
    if valore >= media:
        return 'MEDIA'
    return 'BASSA'


def _riga_valutazione(ws):
    """Riga dell'etichetta 'VALUTAZIONE SU BASE GIORNALIERA', o None."""
    for riga in range(1, ws.max_row + 1):
        for colonna in (COL_CLASSE['HAV'], COL_CLASSE['WBV']):
            testo = _testo(ws.cell(row=riga, column=colonna).value).upper()
            if testo.startswith(ETICHETTA_VALUTAZIONE[:20]):
                return riga
    return None


def _riga_intestazione_tabella(ws, colonna_id, riga_valutazione):
    """
    Riga di intestazione di una delle due tabelle, cercando 'ID' nella colonna
    degli identificativi.

    Il numero di righe dati cambia da scheda a scheda, e i fogli compilati a
    mano prima di VRV hanno la tabella piu' in alto di quella del template:
    conviene trovare l'intestazione invece di darne per scontata la posizione.
    """
    for riga in range(riga_valutazione - 1, 0, -1):
        if _testo(ws.cell(row=riga, column=colonna_id).value).upper() == 'ID':
            return riga
    return RIGA_INIZIO_DATI - 1


def _intestazione_gruppo(ws, riga_limite):
    """
    Codice e descrizione del gruppo omogeneo.

    Prima si prova il blocco a etichette del template (etichetta in colonna C,
    valore in colonna D); se e' vuoto - come nei fogli compilati a mano - si
    prende il primo testo utile fra le righe sopra la tabella.
    """
    intestazione = {}
    for riga, chiave in RIGHE_INTESTAZIONE.items():
        intestazione[chiave] = _testo(
            ws.cell(row=riga, column=COL_VALORE_INTESTAZIONE).value)

    if not intestazione.get('Descrizione_GrOm'):
        ignorare = {'HAV', 'WBV', 'ID'}
        for riga in range(1, max(2, riga_limite)):
            for colonna in range(1, 12):
                testo = _testo(ws.cell(row=riga, column=colonna).value)
                if len(testo) > 3 and testo.upper() not in ignorare:
                    intestazione['Descrizione_GrOm'] = testo
                    break
            if intestazione.get('Descrizione_GrOm'):
                break
    return intestazione


def _leggi_lato(ws, colonne, riga_valutazione, tipo, parametri):
    """Legge una delle due tabelle (HAV o WBV) e il blocco di valutazione."""
    prima_riga = _riga_intestazione_tabella(
        ws, colonne['ID'], riga_valutazione) + 1
    righe = []
    for riga in range(prima_riga, riga_valutazione):
        identificativo = _testo(ws.cell(row=riga, column=colonne['ID']).value)
        if identificativo in ('', '/'):
            continue
        righe.append({
            'id': identificativo,
            'categoria': _testo(ws.cell(row=riga, column=colonne['Categoria']).value),
            'marca': _testo(ws.cell(row=riga, column=colonne['Marca']).value),
            'modello': _testo(ws.cell(row=riga, column=colonne['Modello']).value),
            'aw': _fmt(ws.cell(row=riga, column=colonne['Aw']).value),
            'u': _fmt(ws.cell(row=riga, column=colonne['U']).value, 3),
            'te': _fmt(ws.cell(row=riga, column=colonne['Te']).value, 0),
        })

    a8 = ws.cell(row=riga_valutazione + OFFSET_A8, column=COL_A8[tipo]).value
    ua8 = ws.cell(row=riga_valutazione + OFFSET_A8,
                  column=COL_UA8_AFFIANCO[tipo]).value
    a8max = ws.cell(row=riga_valutazione + OFFSET_A8, column=COL_A8MAX[tipo]).value
    classe = _classe_breve(
        _testo(ws.cell(row=riga_valutazione + OFFSET_CLASSE,
                       column=COL_CLASSE[tipo]).value))
    dy = (ws.cell(row=riga_valutazione + OFFSET_UA8, column=COL_DY_HAV).value
          if tipo == 'HAV' else None)

    if not classe:
        classe = _classe_da_valore(a8max, tipo, parametri)

    return {
        'righe': righe,
        'a8': _fmt(a8),
        'ua8': _fmt(ua8, 3),
        'a8max': _fmt(a8max),
        'dy': _fmt(dy, 1) if tipo == 'HAV' else '',
        'classe': classe or '—',
        'colore': COLORI_CLASSE.get(classe, '#9397ab'),
        'presente': bool(righe) or _numero(a8) is not None,
    }


def leggi_vr_vib(percorso, parametri=None):
    """
    Legge i fogli 'Scheda N' di VR_VIB.xlsx.

    OUTPUT: lista di dict per gruppo omogeneo, con i lati HAV e WBV
    """
    parametri = parametri or {}
    if not percorso or not os.path.exists(percorso):
        return []
    try:
        wb = openpyxl.load_workbook(percorso, data_only=True)
    except Exception:
        return []

    gruppi = []
    for nome_foglio in wb.sheetnames:
        if not nome_foglio.strip().lower().startswith(PREFISSO_SCHEDA):
            continue
        ws = wb[nome_foglio]
        riga_valutazione = _riga_valutazione(ws)
        if riga_valutazione is None:
            continue

        prima_tabella = _riga_intestazione_tabella(
            ws, COL_HAV['ID'], riga_valutazione)
        intestazione = _intestazione_gruppo(ws, prima_tabella)

        codice = intestazione.get('ID_GrOm') or nome_foglio.split(None, 1)[-1].strip()
        gruppi.append({
            'code': codice,
            'nome': intestazione.get('Descrizione_GrOm', ''),
            'reparto': intestazione.get('Descrizione_reparto', ''),
            'HAV': _leggi_lato(ws, COL_HAV, riga_valutazione, 'HAV', parametri),
            'WBV': _leggi_lato(ws, COL_WBV, riga_valutazione, 'WBV', parametri),
        })
    wb.close()
    return gruppi


def leggi_json(cartella_log):
    """Riepilogo scritto dal runner in log/ subito dopo l'analisi."""
    percorso = os.path.join(cartella_log or '', NOME_RIEPILOGO_JSON)
    if not os.path.exists(percorso):
        return None
    try:
        with open(percorso, encoding='utf-8') as f:
            return json.load(f)
    except (json.JSONDecodeError, OSError):
        return None


def sintesi(cartella_log, percorso_vr_vib, parametri=None):
    """
    Struttura per la schermata 'Esposizioni A(8)'.

    INPUT:  cartella_log    - cartella log/ del ramo Vibrazioni (riepilogo json)
            percorso_vr_vib - VR_VIB.xlsx, usato come ripiego
    OUTPUT: {'gruppi': [...], 'conteggi': {'HAV': {...}, 'WBV': {...}},
             'disponibile', 'fonte'}
    """
    parametri = parametri or {}
    dati = leggi_json(cartella_log)
    if dati and dati.get('gruppi'):
        gruppi, fonte = dati['gruppi'], 'json'
    else:
        gruppi, fonte = leggi_vr_vib(percorso_vr_vib, parametri), 'xlsx'

    conteggi = {'HAV': {'ALTA': 0, 'MEDIA': 0, 'BASSA': 0},
                'WBV': {'ALTA': 0, 'MEDIA': 0, 'BASSA': 0}}
    for gruppo in gruppi:
        for tipo in ('HAV', 'WBV'):
            classe = (gruppo.get(tipo) or {}).get('classe', '')
            if classe in conteggi[tipo]:
                conteggi[tipo][classe] += 1

    return {'gruppi': gruppi, 'conteggi': conteggi,
            'disponibile': bool(gruppi), 'fonte': fonte}
