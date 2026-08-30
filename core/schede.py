#! /usr/bin/env python3
"""
Lettura e scrittura di scheda_gruppi_dpi.xlsx.

E' il file condiviso fra le due valutazioni: VRR ne usa le colonne del rumore
(Ti, ID_misura), VRV le colonne HAV e WBV con gli ID delle attrezzature. Ha due
fogli, entrambi con una riga di titolo e le intestazioni sulla riga 2 — per
questo i due backend lo leggono con header=1.

La scrittura avviene con openpyxl sul file originale, riscrivendo le sole celle
dei dati: titolo, intestazioni, larghezze e stile restano intatti. Le righe
aggiunte copiano lo stile dell'ultima riga di dati.
"""

import os
import shutil
from copy import copy

import openpyxl

SCHEDA_DPI = 'Scheda_DPI'
SCHEDA_MANSIONI = 'Scheda_mansioni'

RIGA_TITOLO = 1
RIGA_INTESTAZIONE = 2
RIGA_PRIMO_DATO = 3

COLONNE_DPI = ['codice_DPI', 'Descrizione', 'Marca', 'Modello',
               'SNR', 'H', 'L', 'M', 'Beta']
COLONNE_MANSIONI = ['ID_GrOm', 'Descrizione_GrOm', 'ID_reparto', 'Descrizione_reparto',
                    'Descrizione_compito', 'ID_misura', 'Ti', 'WBV', 'HAV']

# Colonne che devono restare numeriche quando l'utente le compila.
NUMERICHE = {'SNR', 'H', 'L', 'M', 'Beta', 'Ti'}


def _testo(valore):
    """Cella Excel -> stringa. I float interi diventano '240' e non '240.0'."""
    if valore is None:
        return ''
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def _numero(testo):
    """Stringa -> numero se possibile, altrimenti la stringa stessa."""
    testo = (testo or '').strip()
    if testo == '':
        return None
    try:
        valore = float(testo.replace(',', '.'))
    except ValueError:
        return testo
    return int(valore) if valore.is_integer() else valore


def _intestazioni(ws, attese):
    """
    Intestazioni effettive del foglio, completate con quelle attese.

    Se il file e' stato compilato a mano puo' avere colonne in piu' o nomi
    leggermente diversi: si tiene quello che c'e' e si aggiunge il resto.
    """
    lette = []
    for cella in ws[RIGA_INTESTAZIONE]:
        lette.append(_testo(cella.value))
    while lette and lette[-1] == '':
        lette.pop()
    if not lette:
        return list(attese)
    for i, nome in enumerate(attese):
        if i < len(lette) and lette[i] == '':
            lette[i] = nome
    for nome in attese:
        if nome not in lette:
            lette.append(nome)
    return lette


def _ultima_riga_dati(ws, n_colonne):
    """Ultima riga che contiene almeno un valore nelle colonne dei dati."""
    ultima = RIGA_INTESTAZIONE
    for riga in range(RIGA_PRIMO_DATO, ws.max_row + 1):
        for col in range(1, n_colonne + 1):
            if _testo(ws.cell(row=riga, column=col).value) != '':
                ultima = riga
                break
    return ultima


def _leggi_foglio(ws, colonne_attese):
    """
    OUTPUT: {'colonne': [...], 'righe': [[valore, ...], ...]}
    """
    colonne = _intestazioni(ws, colonne_attese)
    ultima = _ultima_riga_dati(ws, len(colonne))
    righe = []
    for riga in range(RIGA_PRIMO_DATO, ultima + 1):
        valori = [_testo(ws.cell(row=riga, column=c).value)
                  for c in range(1, len(colonne) + 1)]
        if any(v != '' for v in valori):
            righe.append(valori)
    return {'colonne': colonne, 'righe': righe}


def leggi(percorso):
    """
    Legge le due tabelle di scheda_gruppi_dpi.xlsx.

    INPUT:  percorso del file
    OUTPUT: {'percorso', 'dpi': {...}, 'mansioni': {...}, 'errore'}
    """
    esito = {'percorso': percorso or '', 'errore': '',
             'dpi': {'colonne': COLONNE_DPI, 'righe': []},
             'mansioni': {'colonne': COLONNE_MANSIONI, 'righe': []}}
    if not percorso or not os.path.exists(percorso):
        esito['errore'] = 'File scheda_gruppi_dpi.xlsx non trovato.'
        return esito

    try:
        wb = openpyxl.load_workbook(percorso, data_only=True)
    except Exception as errore:
        esito['errore'] = f'Impossibile aprire il file: {errore}'
        return esito

    if SCHEDA_DPI in wb.sheetnames:
        esito['dpi'] = _leggi_foglio(wb[SCHEDA_DPI], COLONNE_DPI)
    else:
        esito['errore'] = f'Foglio {SCHEDA_DPI} assente.'
    if SCHEDA_MANSIONI in wb.sheetnames:
        esito['mansioni'] = _leggi_foglio(wb[SCHEDA_MANSIONI], COLONNE_MANSIONI)
    else:
        esito['errore'] = f'Foglio {SCHEDA_MANSIONI} assente.'
    wb.close()
    return esito


def _scrivi_foglio(ws, colonne, righe):
    """
    Riscrive i dati del foglio conservandone titolo, intestazioni e stile.

    Le righe eccedenti vengono svuotate; quelle aggiunte ereditano lo stile
    dell'ultima riga di dati gia' presente.
    """
    n_colonne = len(colonne)
    ultima_esistente = _ultima_riga_dati(ws, n_colonne)
    riga_modello = ultima_esistente if ultima_esistente >= RIGA_PRIMO_DATO else None

    for indice, valori in enumerate(righe):
        riga = RIGA_PRIMO_DATO + indice
        if riga_modello is not None and riga > ultima_esistente:
            for col in range(1, n_colonne + 1):
                origine = ws.cell(row=riga_modello, column=col)
                destinazione = ws.cell(row=riga, column=col)
                destinazione._style = copy(origine._style)
        for col, nome in enumerate(colonne, start=1):
            testo = valori[col - 1] if col - 1 < len(valori) else ''
            cella = ws.cell(row=riga, column=col)
            cella.value = _numero(testo) if nome in NUMERICHE else (testo or None)

    # svuota la coda lasciata dalle righe rimosse
    prima_da_pulire = RIGA_PRIMO_DATO + len(righe)
    for riga in range(prima_da_pulire, ultima_esistente + 1):
        for col in range(1, n_colonne + 1):
            ws.cell(row=riga, column=col).value = None


def salva(percorso, dpi_righe, mansioni_righe, dpi_colonne=None, mansioni_colonne=None):
    """
    Salva le due tabelle sul file originale, previo backup.

    INPUT:  percorso      - scheda_gruppi_dpi.xlsx
            dpi_righe     - lista di liste, nell'ordine delle colonne
            mansioni_righe- idem
    OUTPUT: {'ok', 'errore', 'backup'}
    """
    if not percorso or not os.path.exists(percorso):
        return {'ok': False, 'errore': 'File non trovato.', 'backup': ''}

    backup = percorso + '.bak'
    try:
        if not os.path.exists(backup):
            shutil.copy2(percorso, backup)
        # keep_vba=False: il file e' un .xlsx normale; data_only=False per non
        # sostituire con i valori eventuali formule presenti nel foglio.
        wb = openpyxl.load_workbook(percorso)
        if SCHEDA_DPI in wb.sheetnames:
            _scrivi_foglio(wb[SCHEDA_DPI], dpi_colonne or COLONNE_DPI, dpi_righe)
        if SCHEDA_MANSIONI in wb.sheetnames:
            _scrivi_foglio(wb[SCHEDA_MANSIONI],
                           mansioni_colonne or COLONNE_MANSIONI, mansioni_righe)
        wb.save(percorso)
        wb.close()
    except Exception as errore:
        return {'ok': False, 'errore': str(errore), 'backup': backup}
    return {'ok': True, 'errore': '', 'backup': backup}


def verifica_tempi(mansioni, colonne=None, t0=480.0):
    """
    Somma i Ti per gruppo omogeneo e segnala gli scostamenti da T0.

    analisi_8h di VRR interrompe l'analisi se la somma dei Ti di un gruppo non
    e' esattamente T0: conviene accorgersene qui, con il dettaglio per gruppo,
    invece di leggerlo in un traceback.

    OUTPUT: {'t0', 'gruppi': [{code, nome, tot, delta, oltre}], 'sforati': [...]}
    """
    colonne = colonne or COLONNE_MANSIONI
    try:
        i_id = colonne.index('ID_GrOm')
        i_nome = colonne.index('Descrizione_GrOm')
        i_ti = colonne.index('Ti')
    except ValueError:
        return {'t0': t0, 'gruppi': [], 'sforati': []}

    totali, nomi, ordine = {}, {}, []
    for valori in mansioni:
        if i_id >= len(valori):
            continue
        codice = (valori[i_id] or '').strip()
        if codice == '':
            continue
        if codice not in totali:
            totali[codice] = 0.0
            ordine.append(codice)
        if i_nome < len(valori) and valori[i_nome]:
            nomi.setdefault(codice, valori[i_nome])
        grezzo = valori[i_ti] if i_ti < len(valori) else ''
        try:
            totali[codice] += float(str(grezzo).replace(',', '.')) if grezzo else 0.0
        except ValueError:
            pass

    gruppi = []
    for codice in ordine:
        totale = round(totali[codice], 3)
        delta = round(totale - t0, 3)
        gruppi.append({
            'code': codice,
            'nome': nomi.get(codice, ''),
            'tot': int(totale) if float(totale).is_integer() else totale,
            'delta': f'{"+" if delta > 0 else ""}{int(delta) if float(delta).is_integer() else delta}',
            'oltre': abs(delta) > 1e-6,
        })
    return {'t0': t0, 'gruppi': gruppi,
            'sforati': [g for g in gruppi if g['oltre']]}
