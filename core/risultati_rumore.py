#! /usr/bin/env python3
"""
Lettura dei risultati della valutazione del rischio rumore.

Sono i file che VRR scrive in <Rumore>/output:
    VR8h_riepilogo.xlsx   foglio 'Riepilogo', una riga per gruppo omogeneo
    VR8h_totale.xlsx      un foglio 'Scheda N' per gruppo, con le misure

Qui si rileggono e si trasformano nelle strutture che alimentano le schermate
Superamenti e Misure singole. Si legge anche misure/data/averaged_data.csv, che
e' il punto in cui i valori misurati (U, LeqA, LeqC, Ppeak) sono modificabili a
mano prima di rilanciare l'analisi.
"""

import math
import os

import openpyxl

FOGLIO_RIEPILOGO = 'Riepilogo'
PREFISSO_SCHEDA = 'scheda'

COLONNE_RIEPILOGO = ['ID_GrOm', 'Mansione', 'Lex8h', 'U',
                     'Lex_max', 'L_picco_C', 'classe_rischio']
COLONNE_MISURE = ['ID_GrOm', 'Descrizione_GrOm', 'ID_reparto', 'Descrizione_reparto',
                  'Descrizione_compito', 'ID_misura', 'Ti', 'WBV', 'HAV',
                  'U', 'LeqA', 'LeqC', 'Ppeak']

COLONNE_MEDIE = ['ID', 'U', 'LeqA', 'LeqC', 'Ppeak']
NOME_MEDIE = 'averaged_data.csv'

COLORI_CLASSE = {'BASSA': '#32cd32', 'MEDIA': '#00bfff', 'ALTA': '#b22222'}


def _testo(valore):
    if valore is None:
        return ''
    if isinstance(valore, float):
        if valore.is_integer():
            return str(int(valore))
        return f'{valore:g}'
    return str(valore).strip()


def _numero(valore):
    """Cella -> float, oppure None se non convertibile."""
    if valore is None or valore == '':
        return None
    try:
        return float(str(valore).replace(',', '.'))
    except (TypeError, ValueError):
        return None


def _fmt(valore, decimali=1):
    numero = _numero(valore)
    return '—' if numero is None else f'{numero:.{decimali}f}'


def leggi_riepilogo(percorso):
    """
    OUTPUT: lista di dict con le chiavi di COLONNE_RIEPILOGO (valori testuali)
    """
    if not percorso or not os.path.exists(percorso):
        return []
    try:
        wb = openpyxl.load_workbook(percorso, data_only=True)
    except Exception:
        return []
    ws = wb[FOGLIO_RIEPILOGO] if FOGLIO_RIEPILOGO in wb.sheetnames else wb[wb.sheetnames[0]]

    intestazioni = [_testo(c.value) for c in ws[1]]
    righe = []
    for valori in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in valori):
            continue
        riga = {}
        for i, nome in enumerate(intestazioni):
            if nome:
                riga[nome] = _testo(valori[i]) if i < len(valori) else ''
        if riga.get('ID_GrOm', ''):
            righe.append(riga)
    wb.close()
    return righe


def leggi_totale(percorso):
    """
    Legge i fogli 'Scheda N' di VR8h_totale.xlsx.

    OUTPUT: dizionario ID_GrOm -> lista di dict (una per misura del gruppo)
    """
    if not percorso or not os.path.exists(percorso):
        return {}
    try:
        wb = openpyxl.load_workbook(percorso, data_only=True)
    except Exception:
        return {}

    per_gruppo = {}
    for nome_foglio in wb.sheetnames:
        if not nome_foglio.strip().lower().startswith(PREFISSO_SCHEDA):
            continue
        ws = wb[nome_foglio]
        intestazioni = [_testo(c.value) for c in ws[1]]
        if 'ID_misura' not in intestazioni:
            continue
        for valori in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None, '') for v in valori):
                continue
            riga = {}
            for i, nome in enumerate(intestazioni):
                if nome:
                    riga[nome] = _testo(valori[i]) if i < len(valori) else ''
            codice = riga.get('ID_GrOm', '')
            if codice == '':
                # il nome del foglio e' 'Scheda <ID>': lo si usa come ripiego
                codice = nome_foglio.split(None, 1)[-1].strip()
                riga['ID_GrOm'] = codice
            per_gruppo.setdefault(codice, []).append(riga)
    wb.close()
    return per_gruppo


def _contributi(misure, lex8h, t0):
    """
    Percentuale di energia che ogni misura porta al Lex,8h del gruppo.

    contributo_i = (Ti/T0) * 10^((LeqA_i - Lex8h)/10)
    """
    valore_lex = _numero(lex8h)
    for misura in misure:
        ti = _numero(misura.get('Ti'))
        leqa = _numero(misura.get('LeqA'))
        if valore_lex is None or ti is None or leqa is None or t0 <= 0:
            misura['contributo'] = '—'
            continue
        quota = (ti / t0) * (10 ** ((leqa - valore_lex) / 10.0))
        misura['contributo'] = f'{quota * 100:.1f} %'
    return misure


def sintesi(percorso_riepilogo, percorso_totale, t0=480.0, limite=87.0,
            soglia_media=80.0, soglia_alta=85.0):
    """
    Struttura unica per le schermate Superamenti e Misure singole.

    OUTPUT: {'gruppi': [...], 'conteggi': {...}, 'limite', 'disponibile'}
        ogni gruppo: code, nome, lex, u, lexmax, picco, leqa, ti, classe,
                     oltre_limite, attivita [{id, desc, ti, leqa, leqc,
                     ppeak, contributo, rossa}]
    """
    riepilogo = leggi_riepilogo(percorso_riepilogo)
    totale = leggi_totale(percorso_totale)

    gruppi = []
    for riga in riepilogo:
        codice = riga.get('ID_GrOm', '')
        misure = _contributi(list(totale.get(codice, [])), riga.get('Lex8h'), t0)

        lexmax = _numero(riga.get('Lex_max'))
        classe = (riga.get('classe_rischio') or '').upper()
        if not classe and lexmax is not None:
            classe = ('ALTA' if lexmax >= soglia_alta
                      else 'MEDIA' if lexmax >= soglia_media else 'BASSA')

        leqa_valori = [_numero(m.get('LeqA')) for m in misure]
        leqa_valori = [v for v in leqa_valori if v is not None]
        ti_valori = [_numero(m.get('Ti')) or 0.0 for m in misure]

        attivita = []
        for misura in misure:
            leqa = _numero(misura.get('LeqA'))
            attivita.append({
                'id': misura.get('ID_misura', ''),
                'desc': misura.get('Descrizione_compito', ''),
                'ti': _fmt(misura.get('Ti'), 0),
                'leqa': _fmt(leqa),
                'leqc': _fmt(misura.get('LeqC')),
                'ppeak': _fmt(misura.get('Ppeak')),
                'contributo': misura.get('contributo', '—'),
                # una misura e' "rossa" se da sola supera la soglia di classe alta
                'rossa': leqa is not None and leqa >= soglia_alta,
            })

        gruppi.append({
            'code': codice,
            'nome': riga.get('Mansione', ''),
            'reparto': riga.get('Mansione', ''),
            'lex': _fmt(riga.get('Lex8h')),
            'u': _fmt(riga.get('U')),
            'lexmax': _fmt(riga.get('Lex_max')),
            'picco': _fmt(riga.get('L_picco_C')),
            'leqa': _fmt(max(leqa_valori)) if leqa_valori else '—',
            'ti': f'{sum(ti_valori):g}' if ti_valori else '—',
            'classe': classe or '—',
            'colore': COLORI_CLASSE.get(classe, '#9397ab'),
            'oltre_limite': lexmax is not None and lexmax > limite,
            'n_misure': len(misure),
            'attivita': attivita,
        })

    conteggi = {'ALTA': 0, 'MEDIA': 0, 'BASSA': 0}
    for gruppo in gruppi:
        if gruppo['classe'] in conteggi:
            conteggi[gruppo['classe']] += 1

    return {
        'gruppi': gruppi,
        'conteggi': conteggi,
        'limite': limite,
        'oltre_limite': sum(1 for g in gruppi if g['oltre_limite']),
        'disponibile': bool(gruppi),
    }


def misure_singole(percorso_totale):
    """
    Righe della schermata 'Misure singole', raggruppate per gruppo omogeneo.

    OUTPUT: {'colonne': [...], 'gruppi': [{code, nome, righe: [...]}]}
    """
    totale = leggi_totale(percorso_totale)
    colonne = ['ID_misura', 'Descrizione_compito', 'Ti', 'WBV', 'HAV',
               'U', 'LeqA', 'LeqC', 'Ppeak']
    gruppi = []
    for codice, misure in totale.items():
        gruppi.append({
            'code': codice,
            'nome': misure[0].get('Descrizione_GrOm', '') if misure else '',
            'righe': [{c: m.get(c, '') for c in colonne} for m in misure],
        })
    gruppi.sort(key=lambda g: _chiave_ordinamento(g['code']))
    return {'colonne': colonne, 'gruppi': gruppi}


def _chiave_ordinamento(codice):
    """Ordina '2' prima di '10' e '8.1' prima di '8.31'."""
    try:
        return (0, float(str(codice).replace(',', '.')), '')
    except (TypeError, ValueError):
        return (1, 0.0, str(codice))


# --------------------------------------------------------------------------
# averaged_data.csv: i valori misurati, modificabili prima di rilanciare
# --------------------------------------------------------------------------

def leggi_medie(cartella_data):
    """
    OUTPUT: {'percorso', 'colonne', 'righe': [ {ID, U, LeqA, LeqC, Ppeak, jobName} ]}
    """
    percorso = os.path.join(cartella_data or '', NOME_MEDIE)
    esito = {'percorso': percorso, 'colonne': COLONNE_MEDIE, 'righe': [], 'errore': ''}
    if not os.path.exists(percorso):
        esito['errore'] = 'averaged_data.csv non ancora prodotto.'
        return esito
    try:
        import pandas as pd
        df = pd.read_csv(percorso)
    except Exception as errore:
        esito['errore'] = str(errore)
        return esito

    for _, riga in df.iterrows():
        voce = {'jobName': _testo(riga.get('jobName'))}
        for colonna in COLONNE_MEDIE:
            voce[colonna] = _testo(riga.get(colonna))
        esito['righe'].append(voce)
    return esito


def salva_medie(cartella_data, righe):
    """
    Riscrive averaged_data.csv (e il .xlsx affiancato) con i valori modificati.

    Si aggiornano le sole colonne numeriche U, LeqA, LeqC, Ppeak delle righe
    gia' presenti, riconosciute per ID: il resto del file resta com'e'.
    """
    percorso = os.path.join(cartella_data or '', NOME_MEDIE)
    if not os.path.exists(percorso):
        return {'ok': False, 'errore': 'averaged_data.csv non trovato.'}
    try:
        import pandas as pd
        df = pd.read_csv(percorso)
        per_id = {str(r.get('ID', '')).strip(): r for r in righe}
        for indice, riga in df.iterrows():
            modifica = per_id.get(str(riga.get('ID', '')).strip())
            if not modifica:
                continue
            for colonna in ('U', 'LeqA', 'LeqC', 'Ppeak'):
                valore = _numero(modifica.get(colonna))
                if valore is not None:
                    df.at[indice, colonna] = valore
        df.to_csv(percorso, index=False)
        excel = os.path.splitext(percorso)[0] + '.xlsx'
        try:
            df.to_excel(excel, index=False)
        except Exception:
            pass    # il csv e' la sorgente usata da VRR, l'xlsx e' di comodo
    except Exception as errore:
        return {'ok': False, 'errore': str(errore)}
    return {'ok': True, 'errore': ''}
